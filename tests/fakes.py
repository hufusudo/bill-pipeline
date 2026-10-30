"""造假的微信、支付宝账单：版式照真实账单，数据全是编的。"""
import csv
import io
import json
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace

from openpyxl import Workbook

HEADER = ["交易时间", "交易类型", "交易对方", "商品", "收/支", "金额(元)",
          "支付方式", "当前状态", "交易单号", "商户单号", "备注"]
EXCEL_EPOCH = datetime(1899, 12, 30)


def make_row(n, day, hour, tx_type, who, direction, amount, status, method="零钱"):
    return {
        "交易时间": datetime(2026, 10, day, hour, 0, 0), "交易类型": tx_type, "交易对方": who,
        "商品": "/", "收/支": direction, "金额(元)": amount, "支付方式": method,
        "当前状态": status, "交易单号": f"4200000000202610{n:012d}", "商户单号": "/", "备注": "/",
    }


SAMPLE_ROWS = [
    make_row(1, 1, 12, "商户消费", "食堂一楼", "支出", 12.5, "支付成功"),
    make_row(2, 2, 19, "商户消费", "星光球馆", "支出", 40.0, "支付成功"),
    make_row(3, 3, 8, "扫二维码付款", "张阿姨", "支出", 8.0, "已转账"),
    make_row(4, 3, 21, "转账", "球友A", "支出", 25.0, "对方已收钱"),
    make_row(5, 5, 10, "商户消费", "某电商", "支出", 26.8, "已全额退款"),
    make_row(6, 6, 10, "淘宝平台-退款", "某电商", "收入", 26.8, "已全额退款"),
    make_row(7, 7, 9, "转账", "爸爸", "收入", 1200.0, "已存入零钱"),
    make_row(8, 8, 20, "微信红包", "亲戚", "收入", 88.0, "已存入零钱"),
    make_row(9, 9, 15, "零钱提现", "某银行", "/", 50.5, "提现已到账", method="/"),
]


def _summary_lines(rows):
    totals = {"收入": [0, Decimal(0)], "支出": [0, Decimal(0)], "/": [0, Decimal(0)]}
    for row in rows:
        if row["收/支"] in totals:
            totals[row["收/支"]][0] += 1
            totals[row["收/支"]][1] += Decimal(str(row["金额(元)"]))

    def line(label, key):
        n, amount = totals[key]
        return f"{label}：{n}笔 {amount:.2f}元"

    return {
        "count": f"共{len(rows)}笔记录",
        "income": line("收入", "收入"),
        "expense": line("支出", "支出"),
        "neutral": line("中性交易", "/"),
    }


def write_wechat_xlsx(path, rows=SAMPLE_ROWS, *, title="微信支付账单明细", summary_override=None,
                      time_as="datetime", drop_columns=(), blank_rows_after=0):
    summary = {**_summary_lines(rows), **(summary_override or {})}
    top = [
        title,
        "微信昵称：[测试用户]",
        "起始时间：[2026-10-01 00:00:00] 终止时间：[2026-10-31 23:59:59]",
        "导出类型：[全部账单]",
        "导出时间：[2026-10-31 23:59:59]",
        None,
        summary["count"], summary["income"], summary["expense"], summary["neutral"],
        "注：",
        "1. 充值/提现/理财通购买/零钱通存取/信用卡还款等交易，将计入中性交易",
        "2. 若交易记录明细无有效内容，则代表该时间段内此微信号无交易",
        "3. 本明细仅供个人对账使用",
        "4. 本账单中所有时间均为UTC+08:00时间",
        None,
        "----------------------微信支付账单明细列表--------------------",
    ]
    columns = [name for name in HEADER if name not in drop_columns]

    wb = Workbook()
    ws = wb.active
    for r, text in enumerate(top, start=1):
        if text is not None:
            ws.cell(row=r, column=1, value=text)
    header_row = len(top) + 1
    for c, name in enumerate(columns, start=1):
        ws.cell(row=header_row, column=c, value=name)
    for r, row in enumerate(rows, start=header_row + 1):
        for c, name in enumerate(columns, start=1):
            value = row[name]
            if name == "交易时间" and time_as == "serial":
                value = (value - EXCEL_EPOCH).total_seconds() / 86400
            ws.cell(row=r, column=c, value=value)
    first_blank = header_row + 1 + len(rows)
    for r in range(first_blank, first_blank + blank_rows_after):
        ws.cell(row=r, column=1, value=" ")
    wb.save(path)
    return Path(path)


ALIPAY_HEADER = ["交易时间", "交易分类", "交易对方", "对方账号", "商品说明", "收/支", "金额",
                 "收/付款方式", "交易状态", "交易订单号", "商家订单号", "备注", ""]


def make_alipay_row(n, day, hour, category, who, direction, amount, status="交易成功"):
    return {
        "交易时间": f"2026-10-{day:02d} {hour:02d}:00:00", "交易分类": category, "交易对方": who,
        "对方账号": "/", "商品说明": "测试商品", "收/支": direction, "金额": f"{amount:.2f}",
        "收/付款方式": "余额宝", "交易状态": status,
        # 真实文件里两个单号前后都带制表符
        "交易订单号": f"\t2026100122001{n:015d}\t", "商家订单号": f"\tM{n:010d}\t", "备注": "", "": "",
    }


ALIPAY_ROWS = [
    make_alipay_row(1, 1, 12, "餐饮美食", "某快餐店", "支出", 18.5),
    make_alipay_row(2, 2, 15, "日用百货", "某超市", "支出", 32.0),
    make_alipay_row(3, 3, 20, "文化休闲", "某影院", "支出", 45.0),
    make_alipay_row(4, 4, 9, "投资理财", "某基金", "不计收支", 66.6),
]


def write_alipay_csv(path, rows=ALIPAY_ROWS, *, encoding="gb18030", summary_override=None,
                     with_account_line=True):
    totals = {"收入": [0, Decimal(0)], "支出": [0, Decimal(0)], "不计收支": [0, Decimal(0)]}
    for row in rows:
        if row["收/支"] in totals:
            totals[row["收/支"]][0] += 1
            totals[row["收/支"]][1] += Decimal(row["金额"])
    summary = {
        "count": f"共{len(rows)}笔记录",
        "income": "收入：{}笔 {:.2f}元".format(*totals["收入"]),
        "expense": "支出：{}笔 {:.2f}元".format(*totals["支出"]),
        "neutral": "不计收支：{}笔 {:.2f}元".format(*totals["不计收支"]),
        **(summary_override or {}),
    }
    top = [
        "-" * 84, "导出信息：", "姓名：[测试用户]",
        "支付宝账户：[test@example.com]" if with_account_line else "",
        "起始时间：[2026-10-01 00:00:00]    终止时间：[2026-10-31 23:59:59]",
        "导出交易类型：[全部]", "导出时间：[2026-10-31 23:59:59]",
        summary["count"], summary["income"], summary["expense"], summary["neutral"],
        "", "特别提示：", *[f"{i}.测试提示；" for i in range(1, 9)], "",
        "------------------------支付宝支付科技有限公司  电子客户回单------------------------",
    ]
    out = io.StringIO()
    out.write("\n".join(line or "" for line in top) + "\n")
    writer = csv.writer(out, lineterminator="\n")
    writer.writerow(ALIPAY_HEADER)
    for row in rows:
        writer.writerow([row[name] for name in ALIPAY_HEADER])
    Path(path).write_bytes(out.getvalue().encode(encoding))
    return Path(path)


class FakeLLM:
    """假的模型客户端：按顺序返回预先写好的回复，并记下每次发出去的请求。

    每条回复是 {"content": "..."} 或 {"tool_calls": [(工具名, 参数字典), ...]}。
    """

    def __init__(self, replies):
        self.replies = list(replies)
        self.calls = []
        self.options = []
        self.chat = SimpleNamespace(completions=SimpleNamespace(create=self._create))

    def with_options(self, **kwargs):
        self.options.append(kwargs)
        return self

    def _create(self, **kwargs):
        self.calls.append(kwargs)
        if not self.replies:
            raise RuntimeError("FakeLLM 的回复用完了")
        reply = self.replies.pop(0)
        tool_calls = [SimpleNamespace(id=f"call_{len(self.calls)}_{i}", type="function",
                                      function=SimpleNamespace(name=name, arguments=json.dumps(args, ensure_ascii=False)))
                      for i, (name, args) in enumerate(reply.get("tool_calls", []))] or None
        message = SimpleNamespace(role="assistant", content=reply.get("content"), tool_calls=tool_calls)
        return SimpleNamespace(choices=[SimpleNamespace(message=message)])

    def sent_text(self) -> str:
        """所有请求里发出去的内容拼成一段文字，方便检查有没有泄露个人名字。"""
        return json.dumps(self.calls, ensure_ascii=False, default=str)
